import os
import sys
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_ticket_doc(paper_type="A4"):
    doc = docx.Document()
    section = doc.sections[0]

    if paper_type == "A4":
        p_width = Cm(21.0)
        p_height = Cm(29.7)
        top_m = Cm(1.0)
        bot_m = Cm(1.0)
        left_m = Cm(1.0)
        right_m = Cm(1.0)
    else: # F4 / Folio
        p_width = Cm(21.5)
        p_height = Cm(33.0)
        top_m = Cm(1.5)
        bot_m = Cm(1.5)
        left_m = Cm(1.25)
        right_m = Cm(1.25)

    section.page_width = p_width
    section.page_height = p_height
    section.top_margin = top_m
    section.bottom_margin = bot_m
    section.left_margin = left_m
    section.right_margin = right_m
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)

    # Hapus paragraf default pertama
    if len(doc.paragraphs) > 0:
        p_init = doc.paragraphs[0]._element
        p_init.getparent().remove(p_init)

    logo_path = 'img/header_logos_simplified.png'

    def setup_table_properties(table):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tblPr = table._tbl.tblPr
        tblLayout = parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>')
        tblPr.append(tblLayout)

        tblCellMar = parse_xml(f'''
            <w:tblCellMar {nsdecls("w")}>
                <w:top w:w="0" w:type="dxa"/>
                <w:bottom w:w="0" w:type="dxa"/>
                <w:left w:w="20" w:type="dxa"/>
                <w:right w:w="20" w:type="dxa"/>
            </w:tblCellMar>
        ''')
        tblPr.append(tblCellMar)

        tblBorders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="dashed" w:sz="4" w:space="0" w:color="B0B0B0"/>
                <w:left w:val="dashed" w:sz="4" w:space="0" w:color="B0B0B0"/>
                <w:bottom w:val="dashed" w:sz="4" w:space="0" w:color="B0B0B0"/>
                <w:right w:val="dashed" w:sz="4" w:space="0" w:color="B0B0B0"/>
                <w:insideH w:val="dashed" w:sz="4" w:space="0" w:color="B0B0B0"/>
                <w:insideV w:val="dashed" w:sz="4" w:space="0" w:color="B0B0B0"/>
            </w:tblBorders>
        ''')
        tblPr.append(tblBorders)

    def populate_tickets(table, team_name, prefix, start_num, total_tickets, num_color_rgb, num_font_size=28):
        num = start_num
        for r_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="3100" w:hRule="exact"/>')
            cantSplit = parse_xml(f'<w:cantSplit {nsdecls("w")}/>')
            trPr.append(trHeight)
            trPr.append(cantSplit)

            for c_idx, cell in enumerate(row.cells):
                if num > start_num + total_tickets - 1:
                    break

                cell.width = Cm(4.75)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

                # Elemen 1: Logo Perusahaan (Semen Tonasa & HUT 58 Saja)
                p_logo = cell.paragraphs[0]
                p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_logo.paragraph_format.space_before = Pt(0)
                p_logo.paragraph_format.space_after = Pt(2)
                p_logo.paragraph_format.line_spacing = Pt(12)
                run_logo = p_logo.add_run()
                run_logo.add_picture(logo_path, width=Cm(3.2))

                # Elemen 2: [NAMA TIM MEDIS] ("UTD PANGKEP" atau "PMI MAKASSAR" saja)
                p2 = cell.add_paragraph()
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p2.paragraph_format.space_before = Pt(0)
                p2.paragraph_format.space_after = Pt(2)
                p2.paragraph_format.line_spacing = Pt(10)
                r2 = p2.add_run(team_name)
                r2.font.name = "Arial"
                r2.font.size = Pt(9)
                r2.font.bold = True
                r2.font.color.rgb = RGBColor(33, 37, 41)

                # Elemen 3: NOMOR ANTRIAN (Font 6pt)
                p3 = cell.add_paragraph()
                p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p3.paragraph_format.space_before = Pt(0)
                p3.paragraph_format.space_after = Pt(0)
                p3.paragraph_format.line_spacing = Pt(7)
                r3 = p3.add_run("NOMOR ANTRIAN")
                r3.font.name = "Arial"
                r3.font.size = Pt(6)
                r3.font.color.rgb = RGBColor(108, 117, 125)

                # Elemen 4: [KODE ANTREAN] (Font Bold, Sesuai Tim & Warna)
                p4 = cell.add_paragraph()
                p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p4.paragraph_format.space_before = Pt(0)
                p4.paragraph_format.space_after = Pt(2)
                p4.paragraph_format.line_spacing = Pt(num_font_size)
                r4 = p4.add_run(f"{prefix}-{num:03d}")
                r4.font.name = "Arial Black"
                r4.font.size = Pt(num_font_size)
                r4.font.bold = True
                r4.font.color.rgb = num_color_rgb

                # Elemen 5: Silakan tunggu nomor Anda dipanggil (Font 6.5pt, Italic)
                p5 = cell.add_paragraph()
                p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p5.paragraph_format.space_before = Pt(0)
                p5.paragraph_format.space_after = Pt(1)
                p5.paragraph_format.line_spacing = Pt(7.5)
                r5 = p5.add_run("Silakan tunggu nomor Anda dipanggil")
                r5.font.name = "Arial"
                r5.font.size = Pt(6.5)
                r5.font.italic = True
                r5.font.color.rgb = RGBColor(108, 117, 125)

                # Elemen 6: Terima kasih pahlawan donor! (Font 6.5pt, Bold)
                p6 = cell.add_paragraph()
                p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p6.paragraph_format.space_before = Pt(0)
                p6.paragraph_format.space_after = Pt(0)
                p6.paragraph_format.line_spacing = Pt(7.5)
                r6 = p6.add_run("Terima kasih pahlawan donor!")
                r6.font.name = "Arial"
                r6.font.size = Pt(6.5)
                r6.font.bold = True
                r6.font.color.rgb = RGBColor(33, 37, 41)

                num += 1

    def add_page_break():
        p_break = doc.add_paragraph()
        p_break.paragraph_format.space_before = Pt(0)
        p_break.paragraph_format.space_after = Pt(0)
        p_break.paragraph_format.line_spacing = Pt(1)
        r_brk = p_break.add_run()
        r_brk.font.size = Pt(1)
        r_brk.add_break(WD_BREAK.PAGE)

    # 1. Tim 1: UTD PANGKEP (P-001 s/d P-200, Oranye) -> 50 baris = 10 halaman
    print(f"[{paper_type}] Tim 1: UTD PANGKEP (P-001 s/d P-200, Oranye)...")
    tbl_p = doc.add_table(rows=50, cols=4)
    setup_table_properties(tbl_p)
    populate_tickets(tbl_p, "UTD PANGKEP", "P", 1, 200, RGBColor(230, 81, 0), num_font_size=34)

    # Page Break 1
    add_page_break()

    # 2. Tim 2A: PMI MAKASSAR (MM-001 s/d MM-400, Biru) -> 100 baris = 20 halaman
    print(f"[{paper_type}] Tim 2A: PMI MAKASSAR [MM] (MM-001 s/d MM-400, Biru)...")
    tbl_mm = doc.add_table(rows=100, cols=4)
    setup_table_properties(tbl_mm)
    populate_tickets(tbl_mm, "PMI MAKASSAR", "MM", 1, 400, RGBColor(25, 118, 210), num_font_size=28)

    # Page Break 2
    add_page_break()

    # 3. Tim 2B: PMI MAKASSAR (MO-001 s/d MO-400, Merah) -> 100 baris = 20 halaman
    print(f"[{paper_type}] Tim 2B: PMI MAKASSAR [MO] (MO-001 s/d MO-400, Merah)...")
    tbl_mo = doc.add_table(rows=100, cols=4)
    setup_table_properties(tbl_mo)
    populate_tickets(tbl_mo, "PMI MAKASSAR", "MO", 1, 400, RGBColor(211, 47, 47), num_font_size=28)

    # 4. Lembar 1 Tambahan: Papan Meja / Petunjuk "MO" (Halaman 51, Merah, Ukuran Besar Memenuhi Kertas)
    print(f"[{paper_type}] Menambahkan Lembar 1: Tanda Meja 'MO' (Merah)...")
    sect_mo = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sect_mo.page_width = p_width
    sect_mo.page_height = p_height
    sect_mo.top_margin = top_m
    sect_mo.bottom_margin = bot_m
    sect_mo.left_margin = left_m
    sect_mo.right_margin = right_m
    sect_mo._sectPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

    p_mo = doc.add_paragraph()
    p_mo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mo.paragraph_format.space_before = Pt(0)
    p_mo.paragraph_format.space_after = Pt(0)
    p_mo.paragraph_format.line_spacing = Pt(240)
    r_mo = p_mo.add_run("MO")
    r_mo.font.name = "Arial Black"
    r_mo.font.size = Pt(240)
    r_mo.font.bold = True
    r_mo.font.color.rgb = RGBColor(211, 47, 47)

    # 5. Lembar 2 Tambahan: Papan Meja / Petunjuk "MM" (Halaman 52, Biru, Ukuran Besar Memenuhi Kertas)
    print(f"[{paper_type}] Menambahkan Lembar 2: Tanda Meja 'MM' (Biru)...")
    sect_mm = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sect_mm.page_width = p_width
    sect_mm.page_height = p_height
    sect_mm.top_margin = top_m
    sect_mm.bottom_margin = bot_m
    sect_mm.left_margin = left_m
    sect_mm.right_margin = right_m
    sect_mm._sectPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

    p_mm = doc.add_paragraph()
    p_mm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mm.paragraph_format.space_before = Pt(0)
    p_mm.paragraph_format.space_after = Pt(0)
    p_mm.paragraph_format.line_spacing = Pt(240)
    r_mm = p_mm.add_run("MM")
    r_mm.font.name = "Arial Black"
    r_mm.font.size = Pt(240)
    r_mm.font.bold = True
    r_mm.font.color.rgb = RGBColor(25, 118, 210)

    filename = "Karcis_Antrian_Donor_Darah_STMC.docx" if paper_type == "A4" else "Karcis_Antrian_Donor_Darah_STMC_F4.docx"
    doc.save(filename)
    print(f"[{paper_type}] Sukses menyimpan: {filename}")

if __name__ == '__main__':
    create_ticket_doc("A4")
    create_ticket_doc("F4")
    print("ALL DOCUMENTS GENERATED!")
